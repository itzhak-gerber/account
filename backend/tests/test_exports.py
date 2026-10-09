import io
import zipfile
from typing import Any

from openpyxl import load_workbook

from app.jobs import worker
from tests.fake_idp import FakeIdP
from tests.helpers import app_client
from tests.test_documents import issue, setup
from tests.test_notifications import Jobs, team
from tests.test_payments import get_doc, issued_invoice, make_receipt


def export_jobs(jobs: Jobs) -> list[dict[str, Any]]:
    return [kwargs for name, kwargs in jobs if name == "build_export"]


async def test_full_export_has_every_document_and_the_data(idp: FakeIdP, sent_jobs: Jobs) -> None:
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        invoice = await issued_invoice(owner, bid)
        receipt = await make_receipt(
            owner, bid, "1000", [{"invoice_id": invoice["id"], "amount": "1000"}]
        )
        await issue(owner, bid, receipt["id"])
        sent_jobs.clear()

        started = await owner.post(f"/api/v1/businesses/{bid}/exports")
        assert started.status_code == 202, started.text
        assert started.json()["status"] == "pending"
        busy = await owner.post(f"/api/v1/businesses/{bid}/exports")
        assert busy.status_code == 409
        assert busy.json()["error"]["code"] == "export_in_progress"

        [job] = export_jobs(sent_jobs)
        assert await worker.build_export({}, **job) == "ready"

        [listed] = (await owner.get(f"/api/v1/businesses/{bid}/exports")).json()
        assert listed["status"] == "ready"
        assert listed["documents"] == 2
        assert listed["downloadable"] is True
        download = await owner.get(f"/api/v1/businesses/{bid}/exports/{listed['id']}/download")
        assert download.status_code == 200
        assert download.headers["content-type"] == "application/zip"

        archive = zipfile.ZipFile(io.BytesIO(download.content))
        names = archive.namelist()
        year = invoice["issue_date"][:4]
        assert f"documents/{year}/tax_invoice-000001.pdf" in names
        assert f"documents/{year}/receipt-000001.pdf" in names
        assert {"data.xlsx", "README.txt"} <= set(names)
        assert archive.read(f"documents/{year}/tax_invoice-000001.pdf").startswith(b"%PDF")
        assert "2 מסמכים" in archive.read("README.txt").decode()

        workbook = load_workbook(io.BytesIO(archive.read("data.xlsx")))
        assert workbook.sheetnames == ["מסמכים", "שורות", "תקבולים", "לקוחות", "פריטים"]
        assert workbook["מסמכים"].max_row == 4 + 2  # header rows + two documents
        assert workbook["תקבולים"].max_row >= 5

        # The export carries copies: the single original of each document is not used up.
        assert (await get_doc(owner, bid, invoice["id"]))["original_delivered_at"] is None


async def test_only_owners_admins_and_accountants_export(idp: FakeIdP, sent_jobs: Jobs) -> None:
    async with team(idp, sent_jobs) as (_, owner, bid, member, accountant):
        assert (await member.post(f"/api/v1/businesses/{bid}/exports")).status_code == 403
        assert (await member.get(f"/api/v1/businesses/{bid}/exports")).status_code == 403
        assert (await accountant.post(f"/api/v1/businesses/{bid}/exports")).status_code == 202
        assert len((await owner.get(f"/api/v1/businesses/{bid}/exports")).json()) == 1
