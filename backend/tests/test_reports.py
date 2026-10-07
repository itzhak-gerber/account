from datetime import timedelta
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from app.services.documents import today
from app.services.reports import bucket_for, months_back
from tests.fake_idp import FakeIdP, FakeUser
from tests.helpers import Browser, app_client
from tests.test_documents import CUSTOMER, draft, issue, setup
from tests.test_mcp import call_tool
from tests.test_members import invite_and_join
from tests.test_payments import get_doc, issued_invoice, make_receipt

MIXED_LINES = [
    {"description": "ייעוץ", "quantity": "2", "unit_price": "500"},  # 1000 taxable
    {"description": "שירות לחו״ל", "unit_price": "200", "vat_type": "zero"},
    {"description": "פטור", "unit_price": "300", "vat_type": "exempt"},
]


async def report(owner: Browser, bid: str, name: str, **params: str) -> dict[str, Any]:
    response = await owner.get(f"/api/v1/businesses/{bid}/reports/{name}", params=params)
    assert response.status_code == 200, response.text
    result: dict[str, Any] = response.json()
    return result


async def issued(owner: Browser, bid: str, doc: dict[str, Any]) -> dict[str, Any]:
    response = await issue(owner, bid, doc["id"])
    assert response.status_code == 200, response.text
    result: dict[str, Any] = response.json()
    return result


async def test_income_report_splits_vat_and_subtracts_credit_notes(idp: FakeIdP) -> None:
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        earlier = today() - timedelta(days=40)
        mixed = await issued_invoice(owner, bid, lines=MIXED_LINES, issue_date=earlier.isoformat())
        # Entered with VAT included: 118 is 100 before VAT plus 18 VAT.
        gross = await issued_invoice(
            owner,
            bid,
            lines=[{"description": "מוצר", "unit_price": "118"}],
            prices_include_vat=True,
        )
        credit = (
            await owner.post(f"/api/v1/businesses/{bid}/documents/{gross['id']}/credit-note")
        ).json()
        await issued(owner, bid, credit)
        await draft(owner, bid)  # drafts are not income

        result = await report(
            owner,
            bid,
            "income",
            date_from=earlier.replace(day=1).isoformat(),
            date_to=today().isoformat(),
        )

        assert (mixed["subtotal"], mixed["vat_amount"], mixed["total"]) == (
            "1500.00",
            "180.00",
            "1680.00",
        )
        rows = [
            (d["type"], d["taxable"], d["zero_rated"], d["exempt"], d["vat"], d["total"])
            for d in result["documents"]
        ]
        assert rows == [
            ("tax_invoice", "1000.00", "200.00", "300.00", "180.00", "1680.00"),
            ("tax_invoice", "100.00", "0.00", "0.00", "18.00", "118.00"),
            ("credit_note", "-100.00", "0.00", "0.00", "-18.00", "-118.00"),
        ]
        totals = result["totals"]
        assert (totals["documents"], totals["taxable"], totals["net"], totals["vat"]) == (
            3,
            "1000.00",
            "1500.00",
            "180.00",
        )
        assert result["months"][0]["month"] == earlier.replace(day=1).isoformat()
        assert result["months"][0]["total"] == "1680.00"
        assert sum(m["documents"] for m in result["months"]) == 3


async def test_receipts_report_groups_by_payment_method(idp: FakeIdP) -> None:
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        day = today().isoformat()
        receipt = await owner.post(
            f"/api/v1/businesses/{bid}/documents",
            json={
                "type": "receipt",
                "customer": CUSTOMER,
                "lines": [],
                "payments": [
                    {"method": "cash", "amount": "100", "payment_date": day},
                    {"method": "check", "amount": "200", "payment_date": day},
                    {"method": "cash", "amount": "50", "payment_date": day},
                ],
            },
        )
        await issued(owner, bid, receipt.json())
        invoice_receipt = await draft(
            owner,
            bid,
            type="tax_invoice_receipt",
            lines=[{"description": "שירות", "unit_price": "100"}],
            payments=[{"method": "credit_card", "amount": "118", "payment_date": day}],
        )
        await issued(owner, bid, invoice_receipt)

        result = await report(owner, bid, "receipts")

        assert result["totals"]["by_method"] == {
            "cash": "150.00",
            "check": "200.00",
            "credit_card": "118.00",
        }
        assert (result["totals"]["documents"], result["totals"]["total"]) == (2, "468.00")
        assert [(d["type"], d["total"]) for d in result["documents"]] == [
            ("receipt", "350.00"),
            ("tax_invoice_receipt", "118.00"),
        ]


def test_aging_buckets() -> None:
    assert [bucket_for(d) for d in (-5, 0, 1, 30, 31, 60, 61, 90, 91)] == [
        "current",
        "current",
        "d1_30",
        "d1_30",
        "d31_60",
        "d31_60",
        "d61_90",
        "d61_90",
        "d90_plus",
    ]
    assert months_back(today().replace(month=3, day=15), 11) == today().replace(
        year=today().year - 1, month=4, day=1
    )


async def test_open_balances_by_customer_with_aging(idp: FakeIdP) -> None:
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        old = await issued_invoice(
            owner,
            bid,
            issue_date=(today() - timedelta(days=60)).isoformat(),
            due_date=(today() - timedelta(days=45)).isoformat(),
        )  # 1286.20
        recent = await issued_invoice(owner, bid, customer={"name": "לקוח אחר"})
        paid = await make_receipt(
            owner, bid, "286.20", [{"invoice_id": old["id"], "amount": "286.20"}]
        )
        await issued(owner, bid, paid)

        result = await report(owner, bid, "open-balances")

        assert [
            (d["number"], d["balance"], d["days_overdue"], d["bucket"]) for d in result["documents"]
        ] == [(old["number"], "1000.00", 45, "d31_60"), (recent["number"], "1286.20", 0, "current")]
        assert [(c["customer_name"], c["balance"]) for c in result["customers"]] == [
            ("לקוח אחר", "1286.20"),
            ("לקוח בע״מ", "1000.00"),
        ]
        totals = result["totals"]
        assert (totals["balance"], totals["current"], totals["d31_60"]) == (
            "2286.20",
            "1286.20",
            "1000.00",
        )


async def test_tax_invoice_from_a_paid_proforma_replaces_it(idp: FakeIdP) -> None:
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        proforma = await issued_invoice(owner, bid, type="proforma_invoice")  # 1286.20
        receipt = await make_receipt(
            owner, bid, "500", [{"invoice_id": proforma["id"], "amount": "500"}]
        )
        await issued(owner, bid, receipt)
        converted = await owner.post(
            f"/api/v1/businesses/{bid}/documents/{proforma['id']}/convert",
            json={"type": "tax_invoice"},
        )
        invoice = await issued(owner, bid, converted.json())
        closed = await get_doc(owner, bid, proforma["id"])
        balances = await report(owner, bid, "open-balances")
        open_list = (
            await owner.get(f"/api/v1/businesses/{bid}/documents", params={"open_only": "true"})
        ).json()

        assert (closed["payment_status"], closed["balance_due"]) == ("superseded", "0.00")
        assert (invoice["amount_paid"], invoice["balance_due"]) == ("500.00", "786.20")
        assert [d["id"] for d in balances["documents"]] == [invoice["id"]]
        assert [d["id"] for d in open_list] == [invoice["id"]]


async def test_excel_export_is_rtl_with_numbers_and_no_formulas(idp: FakeIdP) -> None:
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        await issued_invoice(owner, bid, customer={"name": '=HYPERLINK("http://evil","x")'})

        response = await owner.get(f"/api/v1/businesses/{bid}/reports/income.xlsx")
        open_xlsx = await owner.get(f"/api/v1/businesses/{bid}/reports/open_balances.xlsx")

        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert "filename*=UTF-8''" in response.headers["content-disposition"]
        book = load_workbook(BytesIO(response.content))
        assert book.sheetnames == ["סיכום", "מסמכים"]
        details = book["מסמכים"]
        assert details.sheet_view.rightToLeft is True
        header = [c.value for c in details[4]]
        row = {h: c for h, c in zip(header, details[5], strict=True)}
        assert row["סוג מסמך"].value == "חשבונית מס"
        assert row["לקוח"].data_type == "s"  # stored as text, not a formula
        assert row["לקוח"].value.startswith("=HYPERLINK")
        assert row["סה״כ כולל מע״מ"].value == 1286.2
        assert row["סה״כ כולל מע״מ"].number_format == "#,##0.00;-#,##0.00"
        summary = book["סיכום"]
        assert summary.cell(summary.max_row, 1).value == "סה״כ"
        assert open_xlsx.status_code == 200
        assert load_workbook(BytesIO(open_xlsx.content)).sheetnames == ["לפי לקוח", "מסמכים"]


async def test_dashboard_summarises_this_month(idp: FakeIdP) -> None:
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        await issued_invoice(owner, bid)

        result = await report(owner, bid, "dashboard")

        assert result["vat_registered"] is True
        assert (result["income_net"], result["income_vat"]) == ("1090.00", "196.20")
        assert (result["open_balance"], result["open_documents"]) == ("1286.20", 1)
        assert (result["overdue_balance"], result["overdue_documents"]) == ("0.00", 0)
        assert len(result["income_by_month"]) == 12
        assert result["income_by_month"][-1] == {
            "month": today().replace(day=1).isoformat(),
            "amount": "1090.00",
        }


async def test_invalid_period_is_rejected(idp: FakeIdP) -> None:
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        response = await owner.get(
            f"/api/v1/businesses/{bid}/reports/income",
            params={"date_from": "2026-05-01", "date_to": "2026-04-01"},
        )

        assert response.status_code == 422
        assert response.json()["error"]["code"] == "invalid_period"


@pytest.mark.parametrize(("role", "allowed"), [("viewer", False), ("accountant", True)])
async def test_report_permissions(
    idp: FakeIdP, sent_jobs: list[tuple[str, dict[str, Any]]], role: str, allowed: bool
) -> None:
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        user = await invite_and_join(
            client, idp, owner, bid, sent_jobs, FakeUser(email="u@example.com"), role
        )

        for path in ("income", "open-balances", "dashboard", "receipts.xlsx"):
            response = await user.get(f"/api/v1/businesses/{bid}/reports/{path}")
            assert (response.status_code == 200) is allowed, (path, response.status_code)


async def test_mcp_get_report(idp: FakeIdP) -> None:
    owner_user = FakeUser(email="owner@example.com")
    async with app_client() as client:
        owner, bid = await setup(client, idp)
        await issued_invoice(owner, bid)
        token = idp.access_token(owner_user, mfa=True)

        summary = await call_tool(
            client, token, "get_report", {"business_id": bid, "report": "open_balances"}
        )
        detailed = await call_tool(
            client,
            token,
            "get_report",
            {"business_id": bid, "report": "income", "include_documents": True},
        )

        assert summary.get("isError") is not True, summary
        # One tool returns one of three report shapes, so the content is wrapped in "result".
        assert summary["structuredContent"]["result"]["totals"]["balance"] == "1286.20"
        assert summary["structuredContent"]["result"]["documents"] == []
        assert len(detailed["structuredContent"]["result"]["documents"]) == 1
