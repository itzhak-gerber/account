"""Excel (.xlsx) exports of the reports: right-to-left sheets with real numbers and dates, so an
accountant can sum, filter and pivot them."""

from collections.abc import Sequence
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import DocumentType, PaymentMethod
from app.pdf.render import PAYMENT_LABELS
from app.schemas.reports import (
    AgingTotals,
    IncomeDocument,
    IncomeReport,
    IncomeTotals,
    OpenBalancesReport,
    ReceiptsReport,
)
from app.services.document_rules import RULES

CellValue = str | int | Decimal | date | None

MONEY = "#,##0.00;-#,##0.00"
DAY = "dd/mm/yyyy"
MONTH = "mm/yyyy"
BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")


def _title(doc_type: DocumentType) -> str:
    return RULES[doc_type].title_he


def _period(date_from: date, date_to: date) -> str:
    return f"{date_from:%d/%m/%Y} – {date_to:%d/%m/%Y}"


def _sheet(workbook: Workbook, title: str, heading: str, subheading: str) -> Worksheet:
    sheet: Worksheet = workbook.create_sheet(title)
    sheet.sheet_view.rightToLeft = True
    sheet["A1"] = heading
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = subheading
    return sheet


def _table(
    sheet: Worksheet,
    headers: Sequence[str],
    rows: Sequence[Sequence[CellValue]],
    *,
    totals: Sequence[CellValue] | None = None,
    month_column: int | None = None,
    start_row: int = 4,
) -> None:
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(start_row, col, header)
        cell.font, cell.fill = BOLD, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    body = [*rows, totals] if totals is not None else list(rows)
    for r, values in enumerate(body, start=start_row + 1):
        for c, value in enumerate(values, start=1):
            cell = sheet.cell(r, c, value)
            if isinstance(value, str):
                cell.data_type = "s"  # never a formula, whatever a customer name starts with
            elif isinstance(value, Decimal):
                cell.number_format = MONEY
            elif isinstance(value, date):
                cell.number_format = MONTH if c == month_column else DAY
            if totals is not None and values is totals:
                cell.font = BOLD
    for col, header in enumerate(headers, start=1):
        width = max(len(header), *(len(str(row[col - 1] or "")) for row in body)) if body else 10
        sheet.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 40)
    sheet.freeze_panes = f"A{start_row + 1}"


def _save(workbook: Workbook) -> bytes:
    workbook.remove(workbook.worksheets[0])  # the default empty sheet
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def income_xlsx(report: IncomeReport, business_name: str) -> bytes:
    workbook = Workbook()
    period = f"{business_name} · {_period(report.date_from, report.date_to)}"
    amounts = ["עסקאות חייבות", "בשיעור אפס", "פטורות", "סה״כ לפני מע״מ", "מע״מ", "סה״כ כולל מע״מ"]

    def money_cells(row: IncomeTotals | IncomeDocument) -> list[CellValue]:
        return [
            getattr(row, field)
            for field in ("taxable", "zero_rated", "exempt", "net", "vat", "total")
        ]

    summary = _sheet(workbook, "סיכום", "הכנסות ומע״מ", period)
    _table(
        summary,
        ["חודש", "מסמכים", *amounts],
        [[m.month, m.documents, *money_cells(m)] for m in report.months],
        totals=["סה״כ", report.totals.documents, *money_cells(report.totals)],
        month_column=1,
    )
    details = _sheet(workbook, "מסמכים", "מסמכי מס בתקופה", period)
    _table(
        details,
        ["סוג מסמך", "מספר", "תאריך", "לקוח", "מספר עוסק", *amounts],
        [
            [
                _title(d.type),
                d.number,
                d.issue_date,
                d.customer_name,
                d.customer_tax_id,
                *money_cells(d),
            ]
            for d in report.documents
        ],
    )
    return _save(workbook)


def receipts_xlsx(report: ReceiptsReport, business_name: str) -> bytes:
    workbook = Workbook()
    period = f"{business_name} · {_period(report.date_from, report.date_to)}"
    methods = [m for m in PaymentMethod if m in report.totals.by_method]
    method_headers = [PAYMENT_LABELS[m] for m in methods]

    def by_method(values: dict[PaymentMethod, Decimal]) -> list[CellValue]:
        return [values.get(m, Decimal("0.00")) for m in methods]

    summary = _sheet(workbook, "סיכום", "תקבולים", period)
    _table(
        summary,
        ["חודש", "מסמכים", *method_headers, "סה״כ"],
        [[m.month, m.documents, *by_method(m.by_method), m.total] for m in report.months],
        totals=[
            "סה״כ",
            report.totals.documents,
            *by_method(report.totals.by_method),
            report.totals.total,
        ],
        month_column=1,
    )
    details = _sheet(workbook, "מסמכים", "קבלות בתקופה", period)
    _table(
        details,
        ["סוג מסמך", "מספר", "תאריך", "לקוח", *method_headers, "סה״כ"],
        [
            [
                _title(d.type),
                d.number,
                d.issue_date,
                d.customer_name,
                *by_method(d.by_method),
                d.total,
            ]
            for d in report.documents
        ],
    )
    return _save(workbook)


BUCKET_HEADERS = ["טרם הגיע מועד", "1–30 ימים", "31–60 ימים", "61–90 ימים", "מעל 90 ימים"]
BUCKETS = ("current", "d1_30", "d31_60", "d61_90", "d90_plus")


def open_balances_xlsx(report: OpenBalancesReport, business_name: str) -> bytes:
    workbook = Workbook()
    subtitle = f"{business_name} · נכון ל-{report.as_of:%d/%m/%Y}"

    def buckets(row: AgingTotals) -> list[CellValue]:
        return [getattr(row, b) for b in BUCKETS]

    summary = _sheet(workbook, "לפי לקוח", "חובות פתוחים של לקוחות", subtitle)
    _table(
        summary,
        ["לקוח", "מסמכים", *BUCKET_HEADERS, "סה״כ יתרה"],
        [[c.customer_name, c.documents, *buckets(c), c.balance] for c in report.customers],
        totals=["סה״כ", report.totals.documents, *buckets(report.totals), report.totals.balance],
    )
    details = _sheet(workbook, "מסמכים", "מסמכים שלא שולמו במלואם", subtitle)
    _table(
        details,
        ["סוג מסמך", "מספר", "תאריך", "לתשלום עד", "לקוח", "סכום", "שולם", "יתרה", "ימי פיגור"],
        [
            [
                _title(d.type),
                d.number,
                d.issue_date,
                d.due_date,
                d.customer_name,
                d.total,
                d.paid,
                d.balance,
                d.days_overdue,
            ]
            for d in report.documents
        ],
    )
    return _save(workbook)


def export_xlsx(
    business_name: str,
    customers: Sequence[Sequence[CellValue]],
    items: Sequence[Sequence[CellValue]],
    documents: Sequence[Sequence[CellValue]],
    lines: Sequence[Sequence[CellValue]],
    payments: Sequence[Sequence[CellValue]],
) -> bytes:
    """Everything in one workbook, for a full data export."""
    workbook = Workbook()
    _table(
        _sheet(workbook, "מסמכים", "מסמכים שהופקו", business_name),
        [
            "סוג מסמך",
            "מספר",
            "תאריך",
            "לקוח",
            "מספר עוסק",
            "לפני מע״מ",
            "מע״מ",
            "סה״כ",
            "שולם",
            "זוכה",
            "מספר הקצאה",
            "הערות",
        ],
        documents,
    )
    _table(
        _sheet(workbook, "שורות", "שורות המסמכים", business_name),
        [
            "סוג מסמך",
            "מספר",
            "שורה",
            "תיאור",
            "כמות",
            "יחידה",
            "מחיר ליחידה",
            "הנחה %",
            "מע״מ",
            "סה״כ שורה",
        ],
        lines,
    )
    _table(
        _sheet(workbook, "תקבולים", "אמצעי תשלום בקבלות", business_name),
        ["סוג מסמך", "מספר", "אמצעי תשלום", "תאריך", "סכום", "פרטים"],
        payments,
    )
    _table(
        _sheet(workbook, "לקוחות", "לקוחות", business_name),
        ["שם", "מספר עוסק / ת״ז", "דוא״ל", "טלפון", "רחוב", "עיר", "מיקוד", "הערות", "בארכיון"],
        customers,
    )
    _table(
        _sheet(workbook, "פריטים", "פריטים ושירותים", business_name),
        ["שם", "תיאור", "סוג", "מק״ט", "ברקוד", "יחידה", "מחיר", "מע״מ", "בארכיון"],
        items,
    )
    return _save(workbook)
